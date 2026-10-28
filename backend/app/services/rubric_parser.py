"""
Rubric Document Parser Service.
Extracts structured rubric components, marks, and model answers
from uploaded rubric files (PDF, DOCX, TXT) or raw strings.
"""
import os
import re
from typing import Optional


KNOWN_TITLES = [
    "Problem understanding and graph modeling",
    "Choice of algorithm",
    "Initialization of distance matrix",
    "Conversion of train data into graph edges",
    "Floyd-Warshall recurrence relation",
    "Pseudocode correctness",
    "Example execution/output",
    "Time complexity/space complexity",
    "Time complexity / space complexity",
    "Correctness",
    "Completeness",
    "Clarity & Diagrams",
    "Clarity",
    "Approach and Methodology",
    "Optimization and Constraints",
]

DEFAULT_FALLBACK_COMPONENTS = [
    {
        "name": "Correctness & Methodology",
        "marks": 5.0,
        "model_answer": "Complete and theoretically sound solution with valid reasoning and calculations.",
    },
    {
        "name": "Clarity & Completeness",
        "marks": 5.0,
        "model_answer": "Clear presentation, accurate diagrams/pseudocode where applicable, and step-by-step breakdown.",
    },
]


def parse_rubric_text(text: str) -> list[dict]:
    """Parse raw text from a rubric file into structured components."""
    if not text or not text.strip():
        return []

    raw_lines = [l.strip() for l in text.replace("\r", "").split("\n") if l.strip()]
    merged = []

    # Check if numbered items exist
    has_numbered = any(re.match(r"^\d+[\.\)]", l) for l in raw_lines)

    if has_numbered:
        for l in raw_lines:
            if re.match(r"^\d+[\.\)]", l):
                merged.append(l)
            elif (
                merged
                and not l.lower().startswith("total")
                and not l.lower().startswith("rubric")
                and not l.lower().startswith("component")
            ):
                merged[-1] += " " + l
    else:
        # Fallback to lines/bullets/paragraphs
        idx = 1
        for l in raw_lines:
            if l.lower().startswith("total") or l.lower().startswith("rubric:"):
                continue
            clean_l = re.sub(r"^[\*\-•\d\.\)]+\s*", "", l).strip()
            if clean_l and len(clean_l) > 3:
                merged.append(f"{idx}. {clean_l}")
                idx += 1

    components = []
    for item in merged:
        m_num = re.match(r"^(\d+[\.\)])\s*(.*)$", item)
        if not m_num:
            continue
        num_str, rest = m_num.groups()
        m_marks = re.search(r"(\d+(?:\.\d+)?)\s*(?:marks?|pts?)?\s*$", rest, re.IGNORECASE)
        marks = float(m_marks.group(1)) if m_marks else 1.0
        body = rest[:m_marks.start()].strip() if m_marks else rest

        title = body
        content = body

        matched_kt = False
        for kt in KNOWN_TITLES:
            pattern = re.escape(kt).replace(r"\ ", r"\s+").replace("-", ".")
            m_kt = re.search(pattern, body, re.IGNORECASE)
            if m_kt:
                title = body[:m_kt.end()].strip()
                content = body[m_kt.end():].strip()
                if not content:
                    content = title
                matched_kt = True
                break

        if not matched_kt:
            parts = re.split(r"[:\-–—]\s+", body, maxsplit=1)
            if len(parts) == 2:
                title = parts[0].strip()
                content = parts[1].strip()

        components.append({
            "name": f"{num_str} {title}" if not title.startswith(num_str) else title,
            "marks": marks,
            "model_answer": content,
        })

    # If parsing found text but matched zero patterns, provide a single synthesized component
    if not components and raw_lines:
        components.append({
            "name": "1. Assessment Criteria",
            "marks": 10.0,
            "model_answer": "\n".join(raw_lines[:10]),
        })

    return components


def parse_rubric_file(file_path: str) -> list[dict]:
    """Extract and parse structured rubric components from a local file."""
    if not os.path.exists(file_path):
        return []

    ext = os.path.splitext(file_path)[1].lower()
    text = ""

    if ext == ".pdf":
        try:
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(file_path)
            text = "\n".join(p.get_textpage().get_text_range() for p in pdf)
        except Exception as e:
            print(f"[RubricParser] PDF extraction error: {e}")
            return []
    elif ext in (".docx", ".doc"):
        try:
            import docx
            doc = docx.Document(file_path)
            text = "\n".join(p.text for p in doc.paragraphs if p.text)
        except Exception as e:
            print(f"[RubricParser] Word doc extraction error: {e}")
            return []
    else:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        except Exception as e:
            print(f"[RubricParser] Text read error: {e}")
            return []

    return parse_rubric_text(text)
